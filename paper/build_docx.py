"""Build submission Word files from the Markdown manuscript and cover letter."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
IMAGE = re.compile(r"^!\[(.*?)\]\((.+?)\)$")


def set_run_font(run, size, bold=False, italic=False, mono=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    name = "Consolas" if mono else FONT
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline(paragraph, text, size):
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size)
        token = match.group(0)
        run = paragraph.add_run(token[2:-2] if token.startswith("**") else token[1:-1])
        set_run_font(
            run,
            size - 1 if token.startswith("`") else size,
            bold=token.startswith("**"),
            italic=token.startswith("*") and not token.startswith("**"),
            mono=token.startswith("`"),
        )
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size)


def shade_cell(cell):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(document, rows):
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for r_index, row in enumerate(rows):
        for c_index in range(width):
            cell = table.rows[r_index].cells[c_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            value = row[c_index] if c_index < len(row) else ""
            add_inline(paragraph, value, 9)
            if r_index == 0:
                shade_cell(cell)
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def configure(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for style_name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.italic = False


def paragraph(document, text, size=12, after=8, before=0, center=False, hanging=False):
    block = document.add_paragraph()
    fmt = block.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.15
    if center:
        block.alignment = 1
    if hanging:
        fmt.left_indent = Cm(1.0)
        fmt.first_line_indent = Cm(-1.0)
    add_inline(block, text, size)
    return block


def parse_table_row(line):
    parts = [part.strip() for part in line.strip().strip("|").split("|")]
    return parts


def build(source, destination, title, author):
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure(document)
    document.core_properties.title = title
    document.core_properties.author = author
    document.core_properties.subject = "Review Article for the International Journal of Human-Computer Interaction"
    in_references = False
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not re.match(r"^\|\s*-+", lines[index].strip()):
                    rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue
        image = IMAGE.match(line.strip())
        if image:
            caption, relative = image.group(1), image.group(2)
            path = (source.parent / relative).resolve()
            block = document.add_paragraph()
            block.alignment = WD_ALIGN_PARAGRAPH.CENTER
            block.paragraph_format.space_before = Pt(8)
            block.paragraph_format.space_after = Pt(4)
            block.add_run().add_picture(str(path), width=Cm(16.0))
            paragraph(document, caption, size=10, after=12, before=2, center=True)
            index += 1
            continue
        if line.startswith("### "):
            heading = document.add_heading(line[4:], level=2)
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(6)
        elif line.startswith("## "):
            in_references = line[3:].strip() == "References"
            heading = document.add_heading(line[3:], level=1)
            heading.paragraph_format.space_before = Pt(16)
            heading.paragraph_format.space_after = Pt(6)
        elif line.startswith("# "):
            paragraph(document, line[2:], size=16, after=6, center=True)
            for run in document.paragraphs[-1].runs:
                run.bold = True
        else:
            paragraph(
                document,
                line,
                after=6 if in_references else 8,
                hanging=in_references,
            )
        index += 1
    document.save(destination)


if __name__ == "__main__":
    build(
        ROOT / "before-it-acts-manuscript.md",
        ROOT / "before-it-acts-manuscript.docx",
        "Before It Acts: A Scoping Review of Preview and Approval in Agentic AI Interfaces",
        "Tooba Jatoi",
    )
    build(
        ROOT / "before-it-acts-cover-letter.md",
        ROOT / "before-it-acts-cover-letter.docx",
        "Cover letter: Before It Acts",
        "Tooba Jatoi",
    )
